import os
from copy import copy
from pathlib import Path
from shutil import copy2
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn


SHEET_NAME = "Compilado chamados"
TABLE_NAME = "Tabela2"
HEADER_ROW = 4
KEY_COLUMN = "Número do Chamado"

SUGGESTION_COLUMNS = [
    "Intenção Identificada",
    "Causa Identificada",
    "Causa Raiz Padrão",
    "Objeto Operacional Sugerido",
    "Classificação Sugerida",
    "Score Confiança",
]

TRACKING_COLUMNS = [
    "Regra Sugerida",
    "Status da Sugestão",
    "Motivo da Decisão",
    "Ambiguidade",
]

VALIDATION_COLUMNS = [
    "Objeto Operacional Validado",
    "Classificação Validada",
    "Concordância",
]

UPDATE_COLUMNS = (
    SUGGESTION_COLUMNS
    + TRACKING_COLUMNS
)


def _normalize_id(value) -> str:
    """Normaliza o identificador do chamado para comparação."""
    if pd.isna(value):
        return ""

    normalized = str(value).strip()

    if not normalized:
        return ""

    return normalized.split(".")[0]


def _excel_value(value):
    """Converte valores do pandas para valores aceitos pelo openpyxl."""
    if pd.isna(value):
        return None

    # Converte tipos escalares do NumPy para tipos nativos do Python.
    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            pass

    return value


def _copy_row_format_and_formulas(
    worksheet,
    source_row: int,
    target_row: int,
) -> None:
    """
    Copia estilos e fórmulas da linha anterior.

    Valores comuns não são copiados. Apenas fórmulas são traduzidas
    para a nova linha.
    """
    for column in range(
        1,
        worksheet.max_column + 1,
    ):
        source = worksheet.cell(
            source_row,
            column,
        )
        target = worksheet.cell(
            target_row,
            column,
        )

        if source.has_style:
            target._style = copy(source._style)

        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)

        if (
            isinstance(source.value, str)
            and source.value.startswith("=")
        ):
            target.value = Translator(
                source.value,
                origin=source.coordinate,
            ).translate_formula(
                target.coordinate
            )


def _find_last_ticket_row(
    worksheet,
    ticket_column: int,
) -> int:
    """Localiza a última linha que possui Número do Chamado."""
    for row in range(
        worksheet.max_row,
        HEADER_ROW,
        -1,
    ):
        value = worksheet.cell(
            row,
            ticket_column,
        ).value

        if _normalize_id(value):
            return row

    return HEADER_ROW


def _resize_table(
    worksheet,
    last_row: int | None = None,
) -> None:
    """Redimensiona a tabela até a última coluna e linha utilizadas."""
    if TABLE_NAME not in worksheet.tables:
        raise KeyError(
            f"A tabela '{TABLE_NAME}' não foi encontrada "
            f"na aba '{SHEET_NAME}'."
        )

    table = worksheet.tables[TABLE_NAME]

    min_col, min_row, _, current_max_row = (
        range_boundaries(table.ref)
    )

    effective_last_row = (
        last_row
        if last_row is not None
        else current_max_row
    )

    ref = (
        f"{get_column_letter(int(min_col))}{min_row}:"
        f"{get_column_letter(worksheet.max_column)}{effective_last_row}"
    )

    table.ref = ref

    if table.autoFilter is not None:
        table.autoFilter.ref = ref


def _ensure_columns(
    worksheet,
    names: list[str],
) -> None:
    """Cria colunas inexistentes e as inclui na tabela do Excel."""
    table = worksheet.tables[TABLE_NAME]

    existing = {
        str(
            worksheet.cell(
                HEADER_ROW,
                column,
            ).value
        ).strip()
        for column in range(
            1,
            worksheet.max_column + 1,
        )
        if worksheet.cell(
            HEADER_ROW,
            column,
        ).value is not None
    }

    for name in names:
        if name in existing:
            continue

        column = worksheet.max_column + 1

        header = worksheet.cell(
            HEADER_ROW,
            column,
            name,
        )

        previous = worksheet.cell(
            HEADER_ROW,
            column - 1,
        )

        if previous.has_style:
            header._style = copy(previous._style)

        header.font = copy(previous.font)
        header.fill = copy(previous.fill)
        header.border = copy(previous.border)
        header.alignment = copy(previous.alignment)
        header.protection = copy(previous.protection)
        header.number_format = previous.number_format

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = 34

        table.tableColumns.append(
            TableColumn(
                id=len(table.tableColumns) + 1,
                name=name,
            )
        )

        existing.add(name)

    _resize_table(worksheet)


def _build_headers(
    worksheet,
) -> dict[str, int]:
    """Mapeia nome do cabeçalho para número da coluna."""
    headers = {}

    for column in range(
        1,
        worksheet.max_column + 1,
    ):
        value = worksheet.cell(
            HEADER_ROW,
            column,
        ).value

        if value is None:
            continue

        name = str(value).strip()

        if name in headers:
            raise ValueError(
                f"Cabeçalho duplicado na planilha: '{name}'."
            )

        headers[name] = column

    return headers


def _build_row_by_id(
    worksheet,
    ticket_column: int,
) -> dict[str, int]:
    """Mapeia Número do Chamado para sua linha no Excel."""
    row_by_id: dict[str, int] = {}

    duplicated_ids: set[str] = set()

    for row in range(
        HEADER_ROW + 1,
        worksheet.max_row + 1,
    ):
        ticket_id = _normalize_id(
            worksheet.cell(
                row,
                ticket_column,
            ).value
        )

        if not ticket_id:
            continue

        if ticket_id in row_by_id:
            duplicated_ids.add(ticket_id)
            continue

        row_by_id[ticket_id] = row

    if duplicated_ids:
        examples = sorted(duplicated_ids)[:20]

        raise ValueError(
            "Existem Números de Chamado duplicados no compilado: "
            f"{examples}"
        )

    return row_by_id


def _validate_dataframe(
    dataframe: pd.DataFrame,
) -> pd.DataFrame:
    """Valida e normaliza o dataframe recebido pelo exportador."""
    if KEY_COLUMN not in dataframe.columns:
        raise KeyError(
            f"A coluna obrigatória '{KEY_COLUMN}' "
            "não existe no dataframe."
        )

    validated = dataframe.copy()

    validated["_NORMALIZED_TICKET_ID"] = validated[
        KEY_COLUMN
    ].map(_normalize_id)

    validated = validated[
        validated["_NORMALIZED_TICKET_ID"] != ""
    ].copy()

    duplicated = validated[
        validated["_NORMALIZED_TICKET_ID"].duplicated(
            keep=False
        )
    ]

    if not duplicated.empty:
        duplicated_ids = (
            duplicated["_NORMALIZED_TICKET_ID"]
            .drop_duplicates()
            .sort_values()
            .head(20)
            .tolist()
        )

        raise ValueError(
            "O dataframe calculado possui chamados duplicados: "
            f"{duplicated_ids}"
        )

    return validated


def _update_existing_rows(
    worksheet,
    dataframe: pd.DataFrame,
    headers: dict[str, int],
    row_by_id: dict[str, int],
) -> int:
    """Atualiza somente campos calculados pelo AIZEN."""
    updated_count = 0

    available_update_columns = [
        name
        for name in UPDATE_COLUMNS
        if (
            name in dataframe.columns
            and name in headers
        )
    ]

    for _, row in dataframe.iterrows():
        ticket_id = row["_NORMALIZED_TICKET_ID"]
        target_row = row_by_id.get(ticket_id)

        if target_row is None:
            continue

        changed = False

        for name in available_update_columns:
            new_value = _excel_value(
                row.get(name)
            )

            cell = worksheet.cell(
                target_row,
                headers[name],
            )

            if cell.value != new_value:
                cell.value = new_value
                changed = True

        if changed:
            updated_count += 1

    return updated_count


def _write_agreement_formulas(
    worksheet,
    headers: dict[str, int],
    last_row: int,
) -> None:
    """Preenche a fórmula de concordância sem alterar validações."""
    required = {
        "Classificação Sugerida",
        "Classificação Validada",
        "Concordância",
    }

    missing = required.difference(headers)

    if missing:
        raise KeyError(
            "Não foi possível criar a fórmula de concordância. "
            f"Colunas ausentes: {sorted(missing)}"
        )

    suggested_column = get_column_letter(
        headers["Classificação Sugerida"]
    )

    validated_column = get_column_letter(
        headers["Classificação Validada"]
    )

    agreement_column = headers["Concordância"]

    for row in range(
        HEADER_ROW + 1,
        last_row + 1,
    ):
        worksheet.cell(
            row,
            agreement_column,
        ).value = (
            f'=IF(OR('
            f'{suggested_column}{row}="",'
            f'{validated_column}{row}=""'
            f'),"",IF('
            f'{suggested_column}{row}='
            f'{validated_column}{row},'
            f'"SIM","NÃO"))'
        )


def create_output_copy(
    template_path: Path,
    output_path: Path,
) -> Path:
    """Cria uma cópia inicial do template do Plano N2."""
    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.exists():
        raise FileNotFoundError(
            f"Template não encontrado: '{template_path}'."
        )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    copy2(
        src=template_path,
        dst=output_path,
    )

    return output_path


def export_plano_n2(
    dataframe: pd.DataFrame,
    template_path: Path,
    output_path: Path,
) -> tuple[Path, pd.DataFrame]:
    """
    Atualiza as sugestões dos chamados existentes e insere novos chamados.

    Campos manuais e campos de validação não são sobrescritos.
    """
    template_path = Path(template_path)
    output_path = Path(output_path)

    validated_df = _validate_dataframe(
        dataframe
    )

    if not output_path.exists():
        create_output_copy(
            template_path,
            output_path,
        )

    workbook = load_workbook(output_path)

    if SHEET_NAME not in workbook.sheetnames:
        raise KeyError(
            f"A aba '{SHEET_NAME}' não foi encontrada "
            f"em '{output_path}'."
        )

    worksheet = workbook[SHEET_NAME]

    _ensure_columns(
        worksheet,
        (
            SUGGESTION_COLUMNS
            + TRACKING_COLUMNS
            + VALIDATION_COLUMNS
        ),
    )

    headers = _build_headers(
        worksheet
    )

    ticket_column = headers[KEY_COLUMN]

    row_by_id = _build_row_by_id(
        worksheet,
        ticket_column,
    )

    existing_ids = set(row_by_id)

    updated_count = _update_existing_rows(
        worksheet=worksheet,
        dataframe=validated_df,
        headers=headers,
        row_by_id=row_by_id,
    )

    new_rows = validated_df[
        ~validated_df[
            "_NORMALIZED_TICKET_ID"
        ].isin(existing_ids)
    ].copy()

    new_rows.drop(
        columns=["_NORMALIZED_TICKET_ID"],
        inplace=True,
        errors="ignore",
    )

    source_row = _find_last_ticket_row(
        worksheet,
        ticket_column,
    )

    next_row = source_row + 1

    for row_offset, (_, row) in enumerate(
        new_rows.iterrows()
    ):
        target_row = next_row + row_offset

        # Copia estilo e fórmulas antes de inserir os valores.
        _copy_row_format_and_formulas(
            worksheet,
            source_row=source_row,
            target_row=target_row,
        )

        for column_name, column in headers.items():
            if column_name not in new_rows.columns:
                continue

            worksheet.cell(
                target_row,
                column,
            ).value = _excel_value(
                row[column_name]
            )

    last_row = max(
        source_row,
        next_row + len(new_rows) - 1,
    )

    _resize_table(
        worksheet,
        last_row,
    )

    _write_agreement_formulas(
        worksheet,
        headers,
        last_row,
    )

    temporary_path = output_path.with_name(
        f"{output_path.stem}.tmp"
        f"{output_path.suffix}"
    )

    try:
        workbook.save(temporary_path)
        os.replace(
            temporary_path,
            output_path,
        )
    except PermissionError as exc:
        temporary_path.unlink(
            missing_ok=True
        )

        raise PermissionError(
            "O arquivo de saída está aberto ou bloqueado: "
            f"'{output_path}'. Feche o Excel ou Power BI "
            "e execute o pipeline novamente."
        ) from exc
    except Exception:
        temporary_path.unlink(
            missing_ok=True
        )
        raise
    finally:
        workbook.close()

    print(
        "Chamados existentes atualizados: "
        f"{updated_count}"
    )
    print(
        "Novos chamados adicionados: "
        f"{len(new_rows)}"
    )

    return output_path, new_rows