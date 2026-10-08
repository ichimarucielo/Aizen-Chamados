from pathlib import Path
import os

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent

OUTPUT_PATH = (
    BASE_DIR
    / "data"
    / "output"
    / "plano_n2_gerado.xlsx"
)

SHEET_NAME = "Compilado chamados"
HEADER_ROW = 4
TARGET_RULE = "REG_ATENDIMENTO_FINANCEIRO"

COLUMNS_TO_CLEAR = [
    "Intenção Identificada",
    "Causa Identificada",
    "Causa Raiz Padrão",
    "Objeto Operacional Sugerido",
    "Classificação Sugerida",
    "Score Confiança",
    "Regra Sugerida",
    "Status da Sugestão",
    "Motivo da Decisão",
    "Ambiguidade",
]


def main() -> None:
    workbook = load_workbook(
        OUTPUT_PATH
    )

    worksheet = workbook[SHEET_NAME]

    headers = {
        str(
            worksheet.cell(
                HEADER_ROW,
                column,
            ).value
        ).strip(): column
        for column in range(
            1,
            worksheet.max_column + 1,
        )
        if worksheet.cell(
            HEADER_ROW,
            column,
        ).value is not None
    }

    rule_column = headers[
        "Regra Sugerida"
    ]

    cleared = 0

    for row in range(
        HEADER_ROW + 1,
        worksheet.max_row + 1,
    ):
        rule = worksheet.cell(
            row,
            rule_column,
        ).value

        if rule != TARGET_RULE:
            continue

        for column_name in COLUMNS_TO_CLEAR:
            column = headers.get(
                column_name
            )

            if column is not None:
                worksheet.cell(
                    row,
                    column,
                ).value = None

        cleared += 1

    temporary_path = OUTPUT_PATH.with_name(
        f"{OUTPUT_PATH.stem}.tmp"
        f"{OUTPUT_PATH.suffix}"
    )

    try:
        workbook.save(
            temporary_path
        )

        os.replace(
            temporary_path,
            OUTPUT_PATH,
        )
    except Exception:
        temporary_path.unlink(
            missing_ok=True
        )
        raise
    finally:
        workbook.close()

    print(
        "Resultados removidos da regra "
        f"{TARGET_RULE}: {cleared}"
    )


if __name__ == "__main__":
    main()