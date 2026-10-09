"""Reparos pontuais das fórmulas e do layout herdados do Plano N2."""

from openpyxl.utils.cell import get_column_letter, range_boundaries


def repair_legacy_workbook(workbook) -> None:
    """Preserva valores manuais, filtros e fontes das tabelas dinâmicas."""
    if "Compilado chamados" in workbook.sheetnames:
        sheet = workbook["Compilado chamados"]
        for row in range(5, sheet.max_row + 1):
            classification = sheet[f"C{row}"]
            formula = classification.value
            if (
                classification.data_type == "f"
                and formula.startswith("=VLOOKUP(")
                and "'Referencias '!H:I" in formula
            ):
                classification.value = (
                    f'=_xlfn.IFNA({formula[1:]},"Revisão N2 pendente")'
                )

            description = sheet[f"V{row}"]
            formula = description.value
            if (
                description.data_type == "f"
                and "TEXTAFTER(" in formula
                and not formula.startswith(f'=IF(U{row}="",')
            ):
                description.value = (
                    f'=IF(U{row}="","",IFERROR({formula[1:]},U{row}))'
                )

    if "Panorama" not in workbook.sheetnames:
        return
    panorama = workbook["Panorama"]
    if not any(
        panorama[address].data_type == "f"
        and "GETPIVOTDATA(" in panorama[address].value
        and "$B$4" in panorama[address].value
        for address in ("N5", "O5", "P5")
    ):
        return
    for pivot in panorama._pivots:
        if pivot.location.ref != "B4" or pivot.name != "Tabela dinâmica1":
            continue
        # Esse relatório herdado está reduzido a uma célula e seu cache
        # devolve rótulos como se fossem contagens. Os indicadores passam a
        # contar a mesma fonte histórica, sem depender desse cache inválido.
        source = pivot.cache.cacheSource.worksheetSource
        left, top, _, bottom = range_boundaries(source.ref)
        status_index = next(
            i for i, field in enumerate(pivot.cache.cacheFields)
            if field.name == "Status"
        )
        # A contagem direta só é equivalente quando não há filtros ativos.
        for field in pivot.pivotFields:
            if field.axis and any(item.h for item in field.items):
                raise ValueError("Panorama herdado possui filtros: revisar a contagem antes de reparar.")
        column = get_column_letter(left + status_index)
        sheet_name = source.sheet.replace("'", "''")
        status_range = f"'{sheet_name}'!${column}${top + 1}:${column}${bottom}"
        for address, status in (
            ("N5", "PENDENTE"),
            ("O5", "FINALIZADO"),
            ("P5", "EM ANDAMENTO"),
        ):
            cell = panorama[address]
            if cell.data_type == "f" and "GETPIVOTDATA(" in cell.value and "$B$4" in cell.value:
                cell.value = f'=COUNTIF({status_range},"{status}")'
        if panorama["Q4"].value == "A INICIAR" and panorama["Q5"].value == "=SUM(N5:P5)":
            panorama["Q5"] = f'=COUNTIF({status_range},"A INICIAR")'
