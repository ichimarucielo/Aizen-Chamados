from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableColumn

from export import HEADER_ROW, SHEET_NAME, TABLE_NAME, export_plano_n2


def _create_template(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    worksheet.cell(HEADER_ROW, 1, "Número do Chamado")
    worksheet.cell(HEADER_ROW + 1, 1, "BASE-1")

    table = Table(displayName=TABLE_NAME, ref=f"A{HEADER_ROW}:A{HEADER_ROW + 1}")
    table.tableColumns = [TableColumn(id=1, name="Número do Chamado")]
    worksheet.add_table(table)
    workbook.save(path)


def _incoming_ticket() -> pd.DataFrame:
    return pd.DataFrame([{
        "Número do Chamado": "N2-1001",
        "Causa Identificada": "cancelamento da nf",
        "ID Causa Padrão": "CANCELAMENTO_NF",
        "Causa Raiz Padrão": "Cancelamento de Nota Fiscal",
        "Objeto Operacional Sugerido": "Documento Fiscal",
        "Classificação Sugerida": "Cancelamento e Reemissão de Nota Fiscal",
        "Score Confiança": None,
        "Regra Sugerida": "cancelamento_nf_reemissao",
        "Status da Sugestão": "Sugerida",
        "Ambiguidade": False,
        "Objeto Operacional Validado": "",
        "Classificação Validada": "",
        "Concordância": "",
    }])


def test_validation_survives_reexport_and_agreement_is_formula(tmp_path):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    _create_template(template_path)

    export_plano_n2(_incoming_ticket(), template_path, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook[SHEET_NAME]
    headers = {
        worksheet.cell(HEADER_ROW, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    ticket_row = next(
        row
        for row in range(HEADER_ROW + 1, worksheet.max_row + 1)
        if worksheet.cell(row, headers["Número do Chamado"]).value == "N2-1001"
    )
    assert worksheet.cell(ticket_row, headers["Causa Identificada"]).value == "cancelamento da nf"
    assert worksheet.cell(ticket_row, headers["ID Causa Padrão"]).value == "CANCELAMENTO_NF"
    assert worksheet.cell(
        ticket_row,
        headers["Causa Raiz Padrão"],
    ).value == "Cancelamento de Nota Fiscal"
    worksheet.cell(ticket_row, headers["Objeto Operacional Validado"], "Documento Fiscal")
    worksheet.cell(
        ticket_row,
        headers["Classificação Validada"],
        "Cancelamento e Reemissão de Nota Fiscal",
    )
    workbook.save(output_path)

    export_plano_n2(_incoming_ticket(), template_path, output_path)

    workbook = load_workbook(output_path, data_only=False)
    worksheet = workbook[SHEET_NAME]
    assert worksheet.cell(ticket_row, headers["Objeto Operacional Validado"]).value == "Documento Fiscal"
    assert worksheet.cell(
        ticket_row,
        headers["Classificação Validada"],
    ).value == "Cancelamento e Reemissão de Nota Fiscal"
    assert worksheet.cell(ticket_row, headers["Concordância"]).value.startswith("=IF(")


def test_existing_ticket_updates_automatic_fields_preserves_manual_values_and_formulas(tmp_path):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    original = {
        "Número do Chamado": "N2-1001",
        "Status": "Novo",
        "Situação": "PENDENTE",
        "Dias em Aberto": 2,
        "Descrição": "Descrição anterior",
        "Descrição detalhada": "Texto anterior",
        "Assunto": "Assunto anterior",
        "RESPONSÁVEL": "Analista N2",
        "Classificação": "Classificação manual",
        "Causa raiz": "Causa manual",
        "Prioridade": "Alta",
        "Ofensor": "Comercial",
        "Observação/Ação": "Revisado pelo N2",
        "Indicador": "=1+1",
    }
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    for column, (name, value) in enumerate(original.items(), 1):
        worksheet.cell(HEADER_ROW, column, name)
        worksheet.cell(HEADER_ROW + 1, column, value)
    table = Table(displayName=TABLE_NAME, ref=f"A4:N5")
    table.tableColumns = [TableColumn(id=i, name=name) for i, name in enumerate(original, 1)]
    worksheet.add_table(table)
    workbook.save(template_path)
    workbook.close()


    incoming = _incoming_ticket()
    changed = {
        "Status": "Resolvido",
        "Situação": "FINALIZADO",
        "Dias em Aberto": 9,
        "Descrição": "Descrição atualizada",
        "Descrição detalhada": "Texto atualizado",
        "Assunto": "Assunto atualizado",
    }
    for name, value in changed.items():
        incoming[name] = value
    for name in ("RESPONSÁVEL", "Classificação", "Causa raiz", "Prioridade", "Ofensor", "Observação/Ação"):
        incoming[name] = ""

    _, inserted = export_plano_n2(incoming, template_path, output_path)
    assert inserted.empty
    # Backfill contém apenas sugestões; não pode apagar dados operacionais já atualizados.
    sparse = _incoming_ticket()
    sparse["ID Causa Padrão"] = "REENVIO_NF"
    _, inserted_again = export_plano_n2(sparse, template_path, output_path)
    assert inserted_again.empty

    workbook = load_workbook(output_path)
    worksheet = workbook[SHEET_NAME]
    headers = {worksheet.cell(HEADER_ROW, c).value: c for c in range(1, worksheet.max_column + 1)}
    for name, value in original.items():
        assert worksheet.cell(5, headers[name]).value == changed.get(name, value)
    assert worksheet.cell(5, headers["ID Causa Padrão"]).value == "REENVIO_NF"
    assert worksheet.max_row == 5
    workbook.close()


def test_identical_reexport_does_not_report_empty_cells_as_updates(tmp_path, capsys):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"
    _create_template(template_path)
    incoming = _incoming_ticket()
    export_plano_n2(incoming, template_path, output_path)
    capsys.readouterr()

    _, inserted = export_plano_n2(incoming, template_path, output_path)

    assert inserted.empty
    assert "Chamados existentes atualizados: 0" in capsys.readouterr().out
