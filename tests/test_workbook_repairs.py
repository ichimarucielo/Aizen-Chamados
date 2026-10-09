from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from workbook_repairs import repair_legacy_workbook


def test_legacy_formulas_are_guarded_and_manual_values_preserved():
    book = Workbook()
    sheet = book.active
    sheet.title = "Compilado chamados"
    sheet["C5"] = "=VLOOKUP(I:I,'Referencias '!H:I,2,0)"
    sheet["C6"] = "Classe revisada pelo N2"
    sheet["V5"] = '=IFERROR(_xlfn.TEXTAFTER(U5,"marcador"),U5)'
    sheet["V6"] = "Descrição manual"
    sheet["U5"] = None
    repair_legacy_workbook(book)
    assert sheet["C5"].value.endswith(',"Revisão N2 pendente")')
    assert sheet["V5"].value.startswith('=IF(U5="","",IFERROR(')
    assert sheet["V5"].value.endswith(',U5))')
    assert sheet["C6"].value == "Classe revisada pelo N2"
    assert sheet["V6"].value == "Descrição manual"
    first = [sheet["C5"].value, sheet["V5"].value]
    repair_legacy_workbook(book)
    assert first == [sheet["C5"].value, sheet["V5"].value]


def test_real_template_counts_original_source_and_preserves_pivots():
    template = Path(__file__).resolve().parents[1] / "data/input/plano_n2_template.xlsx"
    book = load_workbook(template)
    try:
        pivots = deepcopy(book["Panorama"]._pivots)
        history = deepcopy(book["30-07"]._pivots)
        repair_legacy_workbook(book)
        assert book["Panorama"]["N5"].value == '=COUNTIF(\'Compilado chamados\'!$L$5:$L$283,"PENDENTE")'
        assert book["Panorama"]["O5"].value == '=COUNTIF(\'Compilado chamados\'!$L$5:$L$283,"FINALIZADO")'
        assert book["Panorama"]["Q5"].value == '=COUNTIF(\'Compilado chamados\'!$L$5:$L$283,"A INICIAR")'
        assert book["Panorama"]._pivots == pivots
        assert book["30-07"]._pivots == history
        formula = book["Panorama"]["N5"].value
        repair_legacy_workbook(book)
        assert book["Panorama"]["N5"].value == formula
    finally:
        book.close()


def test_filtered_pivot_is_not_silently_replaced_by_unfiltered_counts():
    template = Path(__file__).resolve().parents[1] / "data/input/plano_n2_template.xlsx"
    book = load_workbook(template)
    try:
        pivot = next(p for p in book["Panorama"]._pivots if p.name == "Tabela dinâmica1")
        pivot.pivotFields[0].items[0].h = True
        with pytest.raises(ValueError, match="possui filtros"):
            repair_legacy_workbook(book)
        assert "GETPIVOTDATA(" in book["Panorama"]["N5"].value
    finally:
        book.close()
